#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================================
 Planilha de resultado no layout do SILOMS
=============================================================================
 Quinze colunas, uma linha por item:

   LOTE | ITEM | REQUISIÇÃO | CNPJ | EMPRESA | QTDE | UND |
   VALOR UNIT LIC | VALOR TOTAL LICI | PRAZO | DESCRIÇÃO |
   SITUAÇÃO | FORNECEDOR | MODELO/VERSAO | MARCA

 Item sem vencedor ocupa uma linha com o aviso no lugar do CNPJ, e nao
 duas, como fazia a ferramenta que serviu de referencia.
=============================================================================
"""

import io
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS = [
    "LOTE", "ITEM", "REQUISIÇÃO", "CNPJ", "EMPRESA", "QTDE", "UND",
    "VALOR UNIT LIC", "VALOR TOTAL LICI", "PRAZO", "DESCRIÇÃO",
    "SITUAÇÃO", "FORNECEDOR", "MODELO/VERSAO", "MARCA",
]

SEM_RESULTADO = "Item deserto/fracassado ou sem retorno na base consultada"
PRAZO_PADRAO = 30


def num(valor, padrao=0.0):
    """Converte para float sem quebrar: número, texto, formato brasileiro, vazio."""
    if valor is None or valor == "":
        return padrao
    if isinstance(valor, (int, float)):
        return float(valor)
    txt = str(valor).strip()
    if "," in txt:
        txt = txt.replace(".", "").replace(",", ".")
    txt = re.sub(r"[^\d.\-]", "", txt)
    try:
        return float(txt)
    except ValueError:
        return padrao


def cnpj_texto(ni):
    """CNPJ/CPF como texto, para o Excel não comer o zero à esquerda."""
    d = re.sub(r"\D", "", str(ni or ""))
    if not d:
        return ""
    return d.zfill(11) if len(d) <= 11 else d.zfill(14)


def _numero_item(item):
    n = item.get("numeroItemCompra")
    try:
        return int(n)
    except (TypeError, ValueError):
        d = re.findall(r"\d+", str(item.get("idCompraItem") or ""))
        try:
            return int(d[-1]) if d else 10 ** 9
        except ValueError:
            return 10 ** 9


def montar(itens, resultados, extras=False, todos_classificados=False):
    """Monta o quadro final cruzando itens com seus vencedores."""
    por_item = {}
    for r in resultados:
        por_item.setdefault(r.get("idCompraItem"), []).append(r)
    for lista in por_item.values():
        lista.sort(key=lambda x: (x.get("ordemClassificacaoSrp") or 9999,
                                  x.get("sequencialResultado") or 9999))

    # sem itens, monta a partir dos proprios resultados
    if not itens and resultados:
        vistos, itens = set(), []
        for r in resultados:
            k = r.get("idCompraItem")
            if k in vistos:
                continue
            vistos.add(k)
            itens.append({"idCompraItem": k,
                          "numeroItemCompra": r.get("numeroItemPncp"),
                          "numeroGrupo": 0, "descricaoResumida": "",
                          "unidadeMedida": ""})

    linhas = []
    for it in sorted(itens, key=_numero_item):
        n = _numero_item(it)
        numero = "" if n == 10 ** 9 else float(n)
        grupo = num(it.get("numeroGrupo"))
        lote = grupo if grupo else ""
        und = (it.get("unidadeMedida") or "") if extras else ""
        desc = (it.get("descricaoResumida") or "") if extras else ""

        vencedores = por_item.get(it.get("idCompraItem"), [])
        if not todos_classificados:
            vencedores = vencedores[:1]

        if not vencedores:
            linhas.append({
                "LOTE": lote, "ITEM": numero, "REQUISIÇÃO": "",
                "CNPJ": SEM_RESULTADO, "EMPRESA": "", "QTDE": 0.0, "UND": und,
                "VALOR UNIT LIC": 0.0, "VALOR TOTAL LICI": 0.0, "PRAZO": None,
                "DESCRIÇÃO": desc, "SITUAÇÃO": "", "FORNECEDOR": "",
                "MODELO/VERSAO": "", "MARCA": "",
            })
            continue

        for v in vencedores:
            qtd = num(v.get("quantidadeHomologada"))
            unit = num(v.get("valorUnitarioHomologado"))
            total = num(v.get("valorTotalHomologado"))
            if total == 0.0 and qtd and unit:
                total = round(qtd * unit, 4)
            linhas.append({
                "LOTE": lote, "ITEM": numero, "REQUISIÇÃO": "",
                "CNPJ": cnpj_texto(v.get("niFornecedor")),
                "EMPRESA": v.get("nomeRazaoSocialFornecedor") or "",
                "QTDE": qtd, "UND": und,
                "VALOR UNIT LIC": unit, "VALOR TOTAL LICI": total,
                "PRAZO": float(PRAZO_PADRAO), "DESCRIÇÃO": desc,
                "SITUAÇÃO": (v.get("situacaoCompraItemResultadoNome") or ""
                             if extras else ""),
                "FORNECEDOR": "", "MODELO/VERSAO": "", "MARCA": "",
            })

    return pd.DataFrame(linhas, columns=COLUNAS)


def gerar_xlsx(df):
    """Devolve o arquivo .xlsx em memória."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet0", index=False)
        ws = writer.sheets["Sheet0"]
        idx = {c: i + 1 for i, c in enumerate(COLUNAS)}

        for i in range(1, len(COLUNAS) + 1):
            cel = ws.cell(row=1, column=i)
            cel.font = Font(bold=True)
            cel.fill = PatternFill("solid", fgColor="DDDDDD")
            cel.alignment = Alignment(horizontal="center", vertical="center")

        for linha in range(2, len(df) + 2):
            for col in ("VALOR UNIT LIC", "VALOR TOTAL LICI"):
                ws.cell(row=linha, column=idx[col]).number_format = "0.0000"
            c = ws.cell(row=linha, column=idx["CNPJ"])
            c.number_format = "@"
            c.alignment = Alignment(horizontal="left")

        larguras = {"LOTE": 8, "ITEM": 8, "REQUISIÇÃO": 12, "CNPJ": 52,
                    "EMPRESA": 45, "QTDE": 10, "UND": 8, "VALOR UNIT LIC": 15,
                    "VALOR TOTAL LICI": 17, "PRAZO": 8, "DESCRIÇÃO": 30,
                    "SITUAÇÃO": 16, "FORNECEDOR": 18, "MODELO/VERSAO": 16,
                    "MARCA": 14}
        for col, larg in larguras.items():
            ws.column_dimensions[get_column_letter(idx[col])].width = larg

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    buffer.seek(0)
    return buffer


def resumo(df):
    com = int((df["CNPJ"] != SEM_RESULTADO).sum())
    sem = int((df["CNPJ"] == SEM_RESULTADO).sum())
    return {
        "itens": int(df["ITEM"].nunique()),
        "com_vencedor": com,
        "sem_vencedor": sem,
        "valor": float(df.loc[df["CNPJ"] != SEM_RESULTADO,
                              "VALOR TOTAL LICI"].sum()),
    }
